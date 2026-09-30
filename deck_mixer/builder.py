"""Deck builders — Pandoro deck mixer.

Each builder turns library cases (or a markdown plan) into a finished .pptx.
Slide layouts live in `slides.py`; colors, fonts and brand assets come from
the `Theme` passed in — see `theme.py`.
"""

from __future__ import annotations

import re
import tempfile
from pathlib import Path

from pptx.util import Inches

from .kb import (
    display_label,
    extract_bullets,
    extract_section,
    get_case_title,
    get_client_display,
    parse_frontmatter,
    read_case,
    redact_client,
    strip_markdown,
)
from .slides import (
    FOOTER_Y, MARGIN, W,
    blank_slide, chip_row, content_header, finish, kpi_tiles,
    new_presentation, rect, slide_bullets, slide_bullets_with_image,
    slide_chart, slide_closing, slide_content, slide_cover_with_image, slide_diagram,
    slide_image_full, slide_kpis, slide_quote, slide_section_divider, slide_statement,
    slide_tech_stack, slide_title_cover, slide_two_col, text_box,
)
from .theme import DEFAULT_THEME, Theme


def _kicker(case: dict) -> str:
    return "  ·  ".join(p for p in (case.get("status", ""), case.get("sector", "")) if p)


def _render_planned(prs, theme: Theme, s, headline: str, bullets: list[str], note: str,
                    eyebrow: str = ""):
    """Render one Glaze-planned slide from its layout + placement."""
    from .glaze import PHOTO

    region = s.placement.region
    photo = s.filled(PHOTO)

    if s.layout == "statement":
        return slide_statement(prs, theme, headline, sub=(bullets[0] if bullets else ""),
                               image_bytes=photo.content if region == "background" and photo else None,
                               note=note)
    if region == "fullbleed" and photo:
        slide = slide_image_full(prs, theme, headline, bullets, photo.content, note=note)
        if eyebrow:
            text_box(slide, MARGIN, Inches(0.4), Inches(6), Inches(0.32),
                     eyebrow.upper(), 10, bold=True, color=theme.accent_dim,
                     font_name=theme.font_body, letter_spacing=1.5)
        return slide
    if region in ("left", "right"):
        return slide_bullets_with_image(prs, theme, headline, bullets,
                                        photo.content if photo else None, note=note,
                                        side=region, eyebrow=eyebrow, hint=s.placement.hint,
                                        lead=s.lead)
    return None


# ---------------------------------------------------------------------------
# Deck builders
# ---------------------------------------------------------------------------

def build_reference_deck(cases: list[dict], kb_path: Path, output_path: Path,
                         template_path: str | None = None,
                         theme: Theme = DEFAULT_THEME) -> str:
    from .visuals import resolve_image
    from .glaze.placeholders import BACKGROUND

    prs = new_presentation(template_path)

    single = len(cases) == 1
    deck_title = get_case_title(cases[0]) if single else "Case References"
    subtitle = (_kicker(cases[0]) if single
                else " | ".join(get_client_display(c) for c in cases[:4]))
    cover_img = resolve_image("artificial intelligence consulting", role=BACKGROUND)
    slide_cover_with_image(prs, theme, deck_title, subtitle, "Reference Case" if single
                           else "Reference Cases", cover_img)

    for case in cases:
        content = read_case(kb_path, case)
        client = get_client_display(case)
        fm = parse_frontmatter(content)

        # Section opener carries a relevant generated/photo background.
        divider_img = resolve_image(case.get("sector", "") or client, role=BACKGROUND)
        sector = case.get("sector", "")
        slide_section_divider(prs, theme, client, image_bytes=divider_img,
                              eyebrow=sector if sector and sector != client else "Reference case")

        slide_content(prs, theme, get_case_title(case), extract_section(content, 1),
                      eyebrow="Summary")

        ctx, challenge = extract_section(content, 2), extract_section(content, 3)
        slide_two_col(prs, theme, "Context & Challenge", "Client context", ctx, "Challenge", challenge)

        combined = extract_section(content, 5) + "\n\n" + extract_section(content, 8)
        slide_bullets(prs, theme, "Approach & Our Role", extract_bullets(combined, 6),
                      eyebrow="Approach")

        slide_content(prs, theme, "Solution", extract_section(content, 6), eyebrow="Solution")

        slide_tech_stack(prs, theme, "Technology & Architecture",
                         extract_section(content, 7), case.get("technologies", []),
                         eyebrow="Tech stack")

        # Native architecture diagram when present in frontmatter
        if fm.get("diagram"):
            slide_diagram(prs, theme, "Solution Architecture", fm["diagram"],
                          caption=fm["diagram"].get("caption", ""),
                          eyebrow="Architecture")

        # KPI tiles when metrics present in frontmatter
        if fm.get("metrics"):
            slide_kpis(prs, theme, "Results in Numbers", fm["metrics"], eyebrow="Impact")

        results = extract_section(content, 10)
        # Native editable chart when a chart spec is present
        if fm.get("chart"):
            slide_chart(prs, theme, "Results & Impact", fm["chart"],
                        bullets=extract_bullets(results, 4),
                        note=extract_section(content, 12), eyebrow="Impact")
        else:
            slide_bullets(prs, theme, "Results & Impact",
                          extract_bullets(results, 5),
                          note=extract_section(content, 12), eyebrow="Impact")

    slide_closing(prs, theme)
    return finish(prs, output_path, deck_title, theme)


# --- Capabilities --------------------------------------------------------------

_TILES_PER_SLIDE = 8        # overview: 4 × 2
_CARDS_PER_SLIDE = 3


def _capability_groups(cases: list[dict], group_by: str) -> list[tuple[str, list[dict]]]:
    """Cases grouped by tag. Tags covering exactly the same cases are merged
    into one group ("Route Optimization · Dispatch Automation") so the deck
    doesn't repeat the same slide under two names."""
    by_tag: dict[str, list[dict]] = {}
    for case in cases:
        for tag in case.get(group_by, []) or []:
            members = by_tag.setdefault(str(tag).strip().lower(), [])
            if case not in members:
                members.append(case)

    merged: dict[tuple, tuple[list[str], list[dict]]] = {}
    for tag, members in by_tag.items():
        key = tuple(sorted(c["case_id"] for c in members))
        merged.setdefault(key, ([], members))[0].append(display_label(tag))

    groups = [(" · ".join(labels), members) for labels, members in merged.values()]
    return sorted(groups, key=lambda g: (-len(g[1]), g[0]))


def _case_metrics(case: dict) -> list[dict]:
    """Frontmatter `metrics` attached by build_capabilities_deck (the catalog doesn't store them)."""
    return [m for m in (case.get("_metrics") or []) if isinstance(m, dict)]


def _card(slide, theme: Theme, case: dict, left, top, width, height, *, wide: bool = False):
    """A case card: client, status/sector kicker, redacted summary, metrics."""
    rect(slide, left, top, width, height, theme.white)
    rect(slide, left, top, width, Inches(0.07), theme.secondary)
    for bx, by, bw, bh in ((left, top, Inches(0.015), height),
                           (left + width - Inches(0.015), top, Inches(0.015), height),
                           (left, top + height - Inches(0.015), width, Inches(0.015))):
        rect(slide, bx, by, bw, bh, theme.hairline)

    pad = Inches(0.28)
    x, w = left + pad, width - pad * 2
    metrics = _case_metrics(case)[:2]
    text_w = w * 0.55 if wide and metrics else w

    text_box(slide, x, top + Inches(0.3), text_w, Inches(0.5), get_client_display(case), 17,
             bold=True, color=theme.primary, font_name=theme.font_head, line_spacing=1.0,
             min_size=12)
    kicker = _kicker(case)
    if kicker:
        sc = theme.accent_text if "production" in case.get("status", "") else theme.mid_gray
        text_box(slide, x, top + Inches(0.8), text_w, Inches(0.3), kicker.upper(), 9,
                 bold=True, color=sc, font_name=theme.font_body, letter_spacing=1.0, min_size=7)

    metrics_h = Inches(1.15) if metrics and not wide else 0
    summary_top = top + Inches(1.2)
    summary = strip_markdown(redact_client(case.get("summary", ""), case))
    text_box(slide, x, summary_top, text_w, top + height - summary_top - metrics_h - Inches(0.2),
             summary, 14 if wide else 13, color=theme.dark_text, font_name=theme.font_body,
             line_spacing=1.25, min_size=10)

    if metrics and wide:
        mx = x + text_w + Inches(0.4)
        kpi_tiles(slide, theme, mx, top + Inches(0.5), x + w - mx, height - Inches(1.0),
                  metrics, value_size=40)
    elif metrics:
        mw = (w - Inches(0.2) * (len(metrics) - 1)) / len(metrics)
        my = top + height - metrics_h - Inches(0.1)
        rect(slide, x, my - Inches(0.08), w, Inches(0.013), theme.hairline)
        for i, m in enumerate(metrics):
            mx = x + i * (mw + Inches(0.2))
            text_box(slide, mx, my, mw, Inches(0.55), str(m.get("value", "")), 26, bold=True,
                     color=theme.secondary_text, font_name=theme.font_head, line_spacing=1.0,
                     min_size=16)
            text_box(slide, mx, my + Inches(0.55), mw, Inches(0.55), str(m.get("label", "")), 10,
                     color=theme.mid_gray, font_name=theme.font_body, line_spacing=1.1, min_size=8)


def _overview_slides(prs, theme: Theme, groups: list[tuple[str, list[dict]]]) -> None:
    cols = 4
    for start in range(0, len(groups), _TILES_PER_SLIDE):
        page = groups[start:start + _TILES_PER_SLIDE]
        slide = blank_slide(prs)
        top = content_header(slide, theme, "What we do", eyebrow="Capabilities")
        gap = Inches(0.25)
        rows = 1 if len(page) <= cols else 2
        tile_w = (W - MARGIN * 2 - gap * (cols - 1)) / cols
        tile_h = min(Inches(2.6), (FOOTER_Y - top - gap * (rows - 1)) / rows)
        for i, (label, members) in enumerate(page):
            x = MARGIN + (i % cols) * (tile_w + gap)
            y = top + (i // cols) * (tile_h + gap)
            rect(slide, x, y, tile_w, tile_h, theme.light_bg)
            rect(slide, x, y, Inches(0.07), tile_h, theme.accent)
            text_box(slide, x + Inches(0.2), y + Inches(0.18), tile_w - Inches(0.35), Inches(0.3),
                     f"{len(members)} CASE{'S' if len(members) > 1 else ''}", 9, bold=True,
                     color=theme.accent_text, font_name=theme.font_body, letter_spacing=1.5)
            text_box(slide, x + Inches(0.2), y + Inches(0.5), tile_w - Inches(0.35),
                     tile_h * 0.45, label, 17, bold=True, color=theme.primary,
                     font_name=theme.font_head, line_spacing=1.05, min_size=11)
            clients = ", ".join(get_client_display(c) for c in members)
            text_box(slide, x + Inches(0.2), y + tile_h * 0.62, tile_w - Inches(0.35),
                     tile_h * 0.34, clients, 11, color=theme.mid_gray,
                     font_name=theme.font_body, line_spacing=1.15, min_size=8)


def build_capabilities_deck(cases: list[dict], kb_path: Path, output_path: Path,
                            group_by: str = "capabilities",
                            template_path: str | None = None,
                            theme: Theme = DEFAULT_THEME) -> str:
    prs = new_presentation(template_path)
    slide_title_cover(prs, theme, f"{theme.company_name}\nCapabilities Overview",
                      f"{len(cases)} reference case{'s' if len(cases) != 1 else ''}",
                      label="Portfolio")

    # Metrics live in the case frontmatter, not the catalog.
    cases = [dict(c, _metrics=parse_frontmatter(read_case(kb_path, c)).get("metrics"))
             for c in cases]
    groups = _capability_groups(cases, group_by)
    _overview_slides(prs, theme, groups)

    for label, members in groups:
        for start in range(0, len(members), _CARDS_PER_SLIDE):
            page = members[start:start + _CARDS_PER_SLIDE]
            slide = blank_slide(prs)
            eyebrow = "Capability" if len(members) <= _CARDS_PER_SLIDE else \
                f"Capability · {start // _CARDS_PER_SLIDE + 1}/{-(-len(members) // _CARDS_PER_SLIDE)}"
            top = content_header(slide, theme, label, eyebrow=eyebrow)
            height = FOOTER_Y - top - Inches(0.1)
            if len(page) == 1:
                _card(slide, theme, page[0], MARGIN, top, W - MARGIN * 2, height, wide=True)
                continue
            gap = Inches(0.3)
            card_w = (W - MARGIN * 2 - gap * (_CARDS_PER_SLIDE - 1)) / _CARDS_PER_SLIDE
            for i, case in enumerate(page):
                _card(slide, theme, case, MARGIN + i * (card_w + gap), top, card_w, height)

    slide_closing(prs, theme)
    return finish(prs, output_path, f"{theme.company_name} Capabilities Overview", theme)


# --- Tender ----------------------------------------------------------------------

# Case-evidence slides are built from real case data (summary + results),
# not invented quotes or charts — so only offer layouts we can honestly fill.
_TENDER_LAYOUTS = ("statement", "bullets", "image_left", "image_right", "image_full")


def build_tender_deck(cases: list[dict], kb_path: Path, output_path: Path,
                      brief: str = "", template_path: str | None = None,
                      theme: Theme = DEFAULT_THEME) -> str:
    from .glaze import glaze
    from .visuals import resolve_image
    from .glaze.placeholders import BACKGROUND

    prs = new_presentation(template_path)

    cover_img = resolve_image("AI strategy consulting", role=BACKGROUND)
    slide_cover_with_image(prs, theme, f"{theme.company_name}\nReference Cases",
                           brief or "Tender Support", "Tender Reference", cover_img)

    if brief:
        slide_content(prs, theme, "Tender Context", brief, eyebrow="Brief")

    # Build one storyline section per case, so the art-director planner (Glaze)
    # can vary layout/photo treatment across the whole deck instead of
    # repeating the same template on every case slide.
    sections: list[tuple[str, str]] = []
    case_meta: list[dict] = []
    for case in cases:
        content = read_case(kb_path, case)
        summary = extract_section(content, 1)
        results = extract_section(content, 10)
        body = f"{strip_markdown(summary)}\n\n{strip_markdown(results)}"
        sections.append((get_case_title(case), body))
        case_meta.append({
            "client": get_client_display(case),
            "kicker": _kicker(case),
            "tender": extract_section(content, 12),
            "chips": case.get("technologies", [])[:6],
        })

    slides = glaze(sections, layouts=_TENDER_LAYOUTS, company=theme.company_name, theme=theme)

    for (heading, body), s, meta in zip(sections, slides, case_meta):
        headline = s.headline or heading
        bullets = s.bullets or extract_bullets(body, 5)
        note = "\n\n".join(p for p in (
            meta["kicker"], meta["tender"] or s.notes,
            "Tech: " + ", ".join(meta["chips"]) if meta["chips"] else "") if p)

        slide = _render_planned(prs, theme, s, headline, bullets, note, eyebrow=meta["client"])
        if slide is None:
            chips_h = Inches(0.55) if meta["chips"] else 0
            slide = slide_bullets(prs, theme, headline, bullets, note=note,
                                  eyebrow=meta["client"], reserve_bottom=chips_h)
            if meta["chips"]:
                chip_row(slide, theme, MARGIN, FOOTER_Y - Inches(0.45), meta["chips"],
                         font_size=10, right=W - Inches(2.3), max_rows=1)

    slide_closing(prs, theme)
    return finish(prs, output_path, f"{theme.company_name} Reference Cases", theme)


# ---------------------------------------------------------------------------
# Plan / proposal deck builder
# ---------------------------------------------------------------------------

_PHASE_KEYWORDS = ("fase ", "phase ", "stap ", "etappe ")


def _is_phase_heading(heading: str) -> bool:
    return any(heading.lower().startswith(k) for k in _PHASE_KEYWORDS)


def build_plan_deck(
    title: str,
    content_md: str,
    subtitle: str = "",
    label: str = "Plan of Approach",
    output_path: Path = Path(tempfile.gettempdir()) / "plan.pptx",
    template_path: str | None = None,
    enrich: bool = True,
    company: str = "",
    theme: Theme = DEFAULT_THEME,
) -> str:
    """Build a plan/proposal deck via the Glaze engine.

    Glaze runs three phases — Recipe (plan slides + define typed placeholders),
    Proof (evaluate which placeholders earn their place), and Pour (call the
    sources to fill them). With an LLM key configured the Recipe phase uses it;
    otherwise it falls back to layout heuristics. Set enrich=False to stop
    after Proof (no images fetched; image slots become placeholders).
    """
    from .glaze import glaze, DIAGRAM, CHART, KPI

    prs = new_presentation(template_path)

    parts = re.split(r"^## (.+)$", content_md, flags=re.MULTILINE)
    sections = [
        (parts[i].strip(), parts[i + 1].strip() if i + 1 < len(parts) else "")
        for i in range(1, len(parts) - 1, 2)
    ]

    # Glaze: Plan (demand + describe) -> Fetch (get images) -> Place.
    slides = glaze(sections, company=company, fill=enrich, theme=theme)

    # Cover — give the title slide a dedicated hero image (the cover always
    # warrants one). Falls back to the geometric cover when none comes back.
    cover_img = None
    if enrich:
        from .visuals import resolve_image
        from .glaze.placeholders import BACKGROUND
        first = sections[0][0] if sections else (title.replace("\n", " "))
        cover_img = resolve_image(
            first, role=BACKGROUND,
            brief=f"A premium abstract hero image for a {theme.company_name} proposal titled "
                  f"“{title.replace(chr(10), ' ')}” about {first}; {theme.palette_description}, "
                  f"modern, optimistic, lots of negative space.")
    slide_cover_with_image(prs, theme, title, subtitle, label, cover_img)

    for (heading, body), s in zip(sections, slides):
        if _is_phase_heading(heading):
            slide_section_divider(prs, theme, heading, eyebrow="Phase")
            if s.layout == "statement":     # the divider already is the big moment
                s.layout = "bullets"

        headline = s.headline or heading
        bullets = s.bullets or extract_bullets(body, 5)
        note = s.notes or strip_markdown(body)

        region = s.placement.region
        diagram = s.filled(DIAGRAM)
        chart = s.filled(CHART)
        kpi = s.filled(KPI)

        if s.layout == "quote":
            slide_quote(prs, theme, headline, attribution=company, note=note)
        elif _render_planned(prs, theme, s, headline, bullets, note):
            pass
        elif region == "full" and diagram:
            slide_diagram(prs, theme, headline, diagram.content, note=note)
        elif region == "full" and chart:
            slide_chart(prs, theme, headline, chart.content, bullets=bullets, note=note)
        elif region == "full" and kpi:
            slide_kpis(prs, theme, headline, kpi.content.get("metrics", []), note=note)
        else:  # clean information slide
            slide_bullets(prs, theme, headline, bullets, note=note, lead=s.lead)

    slide_closing(prs, theme)
    return finish(prs, output_path, title, theme)
