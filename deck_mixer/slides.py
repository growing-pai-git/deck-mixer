"""Slide layouts for the deck mixer.

Visual design: dark/accent/white, geometric cover accents, styled bullets,
split section dividers, two-zone image+text layout. All colors, fonts and
brand assets come from the `Theme` passed into each layout — see `theme.py`.

Every body text box is fitted to its box (see `textfit.py`): the font
shrinks within a readable range and, only if that isn't enough, trailing
sentences are dropped. Nothing is allowed to run past the footer.
"""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .kb import strip_markdown
from .textfit import BULLET_GAP_EM, fit_bullets, fit_text, text_height
from .theme import Theme

# ---------------------------------------------------------------------------
# Slide geometry: widescreen 13.33 × 7.5"
# ---------------------------------------------------------------------------
W          = Inches(13.33)
H          = Inches(7.5)
MARGIN     = Inches(0.55)
STRIP_W    = Inches(0.09)   # left accent strip on the image cover
FOOTER_Y   = H - Inches(0.7)  # body text must end above the footer brand mark

CREDIT = "Made with Deck Mixer by Growing pAI"


# ---------------------------------------------------------------------------
# Presentation
# ---------------------------------------------------------------------------

def new_presentation(template_path: str | None = None) -> Presentation:
    if not template_path:
        prs = Presentation()
        prs.slide_width = W
        prs.slide_height = H
        return prs

    p = Path(template_path).expanduser()
    if not p.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")
    prs = Presentation(str(p))
    ratio = prs.slide_width / prs.slide_height
    if abs(ratio - W / H) > 0.01:
        raise ValueError(
            f"Template {p.name} is not 16:9 widescreen "
            f"({prs.slide_width / 914400:.2f} × {prs.slide_height / 914400:.2f} in). "
            "Deck Mixer lays slides out for 16:9 — resize the template in "
            "PowerPoint (Design → Slide Size → Widescreen) and try again.")
    if len(prs.slides):
        raise ValueError(
            f"Template {p.name} already contains {len(prs.slides)} slide(s). "
            "Save it with no slides (or as a .potx) so they don't end up in every deck.")
    prs.slide_width, prs.slide_height = W, H
    return prs


def _blank_layout(prs: Presentation):
    """The layout with the fewest placeholders — "Blank" in the default
    template, but corporate templates order their layouts differently."""
    for layout in prs.slide_layouts:
        if layout.name.strip().lower() == "blank":
            return layout
    return min(prs.slide_layouts, key=lambda lay: len(lay.placeholders))


def blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(_blank_layout(prs))
    for ph in list(slide.placeholders):         # never leave "Click to add title" behind
        ph._element.getparent().remove(ph._element)
    return slide


def finish(prs: Presentation, output_path: Path, title: str, theme: Theme) -> str:
    """Stamp document properties and save. Without this python-pptx leaves
    its own template's author ("Steve Canny") and 2013 dates on every deck."""
    now = datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)
    cp = prs.core_properties
    cp.title = " ".join(title.split())
    cp.author = theme.company_name
    cp.last_modified_by = theme.company_name
    cp.created = now
    cp.modified = now
    cp.revision = 1
    cp.comments = CREDIT
    cp.subject = ""
    cp.keywords = ""
    prs.save(str(output_path))
    return str(output_path)


def set_notes(slide, *parts: str) -> None:
    text = "\n\n".join(p.strip() for p in parts if p and p.strip())
    if text:
        slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------------------
# Primitive helpers
# ---------------------------------------------------------------------------

def rect(slide, left, top, width, height, fill: RGBColor, rotation: float = 0,
         shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, int(left), int(top), int(width), int(height))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if rotation:
        s.rotation = rotation
    return s


def scrim(slide, left, top, width, height, fill: RGBColor, alpha_pct: int = 60):
    """Semi-transparent colour overlay (alpha_pct = opacity %)."""
    s = rect(slide, left, top, width, height, fill)
    srgb = s.fill.fore_color._xFill.find(qn("a:srgbClr"))
    if srgb is not None:
        srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha_pct * 1000))}))
    return s


def text_box(
    slide, left, top, width, height, text: str,
    font_size: float = 14,
    bold: bool = False,
    color: RGBColor = RGBColor(0x1A, 0x22, 0x30),
    align: PP_ALIGN = PP_ALIGN.LEFT,
    italic: bool = False,
    line_spacing: float = 1.18,
    font_name: str = "Calibri",
    letter_spacing: float | None = None,
    min_size: float | None = None,
    anchor: MSO_ANCHOR | None = None,
):
    """A text box; newlines in `text` become paragraphs.

    With `min_size`, the text is fitted to the box: the font shrinks from
    `font_size` towards `min_size`, then trailing sentences are dropped.
    """
    if min_size is not None:
        font_size, text = fit_text(text, width, height, font_size, min_size,
                                   line_spacing=line_spacing, bold=bold)

    tb = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    if min_size is not None:
        # Safety net for later hand edits: PowerPoint shrinks on overflow.
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = font_name
        if letter_spacing is not None:
            # Tracking in 1/100 pt, applied via direct XML on the run
            run._r.get_or_add_rPr().set("spc", str(int(letter_spacing * 100)))
    return tb


def body_text(slide, theme: Theme, left, top, width, height, text: str,
              size: float = 15, min_size: float = 11):
    """Prose body copy: markdown stripped, paragraphs kept, fitted to the box."""
    paragraphs = [p for p in strip_markdown(text).split("\n\n") if p.strip()]
    return text_box(slide, left, top, width, height, "\n".join(paragraphs), size,
                    color=theme.dark_text, font_name=theme.font_body,
                    line_spacing=1.3, min_size=min_size)


def _cover_crop(image_bytes: bytes, width, height) -> bytes:
    """Center-crop to the target box's aspect ratio (a 'cover' fit), so
    embedding fills the box without stretch-distortion. Image generators
    (e.g. Gemini Flash) often return a fixed square regardless of the
    requested orientation, so this runs on every embed."""
    from PIL import Image
    im = Image.open(BytesIO(image_bytes))
    target_ratio = width / height
    src_ratio = im.width / im.height
    if abs(target_ratio - src_ratio) < 0.02:
        return image_bytes
    if src_ratio > target_ratio:                    # source too wide: crop sides
        new_w = round(im.height * target_ratio)
        x = (im.width - new_w) // 2
        im = im.crop((x, 0, x + new_w, im.height))
    else:                                            # source too tall: crop top/bottom
        new_h = round(im.width / target_ratio)
        y = (im.height - new_h) // 2
        im = im.crop((0, y, im.width, y + new_h))
    buf = BytesIO()
    im.convert("RGB").save(buf, format="JPEG", quality=92)
    return buf.getvalue()


def embed_image(slide, image_bytes: bytes, left, top, width, height, theme: Theme):
    try:
        image_bytes = _cover_crop(image_bytes, width, height)
        slide.shapes.add_picture(BytesIO(image_bytes), int(left), int(top),
                                 int(width), int(height))
    except Exception:
        rect(slide, left, top, width, height, theme.light_bg)


PLACEHOLDER_NAME = "Image placeholder"


def image_placeholder(slide, theme: Theme, left, top, width, height, hint: str = ""):
    """A clearly-marked, on-brand empty image slot, grouped so it can be
    selected and deleted in one click. Used whenever no image source is
    configured — the deck never ships a random stock photo the user didn't
    choose."""
    group = slide.shapes.add_group_shape()
    group.name = PLACEHOLDER_NAME

    box = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(left), int(top), int(width), int(height))
    box.fill.solid()
    box.fill.fore_color.rgb = theme.light_bg
    box.line.color.rgb = theme.hairline
    box.line.width = Pt(1.25)
    box.line.dash_style = 4                      # MSO_LINE_DASH_STYLE.DASH
    box.shadow.inherit = False

    # A small "picture" glyph: frame, sun, mountain.
    gw, gh = Inches(0.9), Inches(0.66)
    gx = left + (width - gw) / 2
    gy = top + height / 2 - gh - Inches(0.12)
    frame = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(gx), int(gy), int(gw), int(gh))
    frame.fill.background()
    frame.line.color.rgb = theme.mid_gray
    frame.line.width = Pt(1.5)
    frame.shadow.inherit = False
    for shape, x, y, w, h in (
        (MSO_SHAPE.OVAL, gx + gw * 0.62, gy + gh * 0.18, gw * 0.16, gw * 0.16),
        (MSO_SHAPE.ISOSCELES_TRIANGLE, gx + gw * 0.12, gy + gh * 0.38, gw * 0.5, gh * 0.5),
    ):
        s = group.shapes.add_shape(shape, int(x), int(y), int(w), int(h))
        s.fill.solid()
        s.fill.fore_color.rgb = theme.mid_gray
        s.line.fill.background()
        s.shadow.inherit = False

    label = group.shapes.add_textbox(int(left + Inches(0.3)), int(top + height / 2 + Inches(0.02)),
                                     int(width - Inches(0.6)), int(Inches(0.9)))
    tf = label.text_frame
    tf.word_wrap = True
    for i, (text, size, bold) in enumerate((("IMAGE PLACEHOLDER", 10, True),
                                            (hint, 11, False))):
        if not text:
            continue
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text if i == 0 else f"Suggested: {text}"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = not bold
        r.font.color.rgb = theme.mid_gray
        r.font.name = theme.font_body
        if bold:
            r._r.get_or_add_rPr().set("spc", "150")

    group._element.nvGrpSpPr.cNvPr.set(
        "descr", f"Replace with your own image{f' ({hint})' if hint else ''}.")
    return group


def placeholder_note(hint: str) -> str:
    what = f" — suggested: {hint}" if hint else ""
    return (f"Image placeholder{what}. Replace it with your own image: select "
            "the grey box, delete it, then Insert → Picture.")


def favicon(slide, theme: Theme, left, top, size=Inches(0.34)):
    """Place the compact brand mark — the recurring visual cue used on most
    slides. Falls back to a small secondary-colour square if no asset is set."""
    if theme.favicon_path and theme.favicon_path.exists():
        try:
            slide.shapes.add_picture(str(theme.favicon_path), int(left), int(top),
                                     width=int(size), height=int(size))
            return
        except Exception:
            pass
    rect(slide, left, top, size, size, theme.secondary)


def logo(slide, theme: Theme, *, dark_bg: bool, left, top, height=Inches(0.42),
         width=Inches(2.4), align: PP_ALIGN = PP_ALIGN.LEFT):
    """Place the brand logo (reserved for brand slides — cover, closing).

    On dark slides, if the logo isn't transparent-safe on the primary colour
    it sits on a clean white rounded plate. Falls back to a text wordmark if
    no logo asset is configured. `width`/`align` position the wordmark
    fallback; an image logo is placed at `left` at its natural aspect ratio.
    """
    if theme.logo_path and theme.logo_path.exists():
        try:
            from PIL import Image
            with Image.open(theme.logo_path) as im:
                ar = im.width / im.height
            w = int(height * ar)
            if align == PP_ALIGN.CENTER:
                left = left + (width - w) / 2
            if dark_bg:
                pad = Inches(0.18)
                plate = rect(slide, left - pad, top - pad, w + pad * 2, height + pad * 2,
                             theme.white, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                plate.shadow.inherit = False
            slide.shapes.add_picture(str(theme.logo_path), int(left), int(top),
                                     width=w, height=int(height))
            return
        except Exception:
            pass
    color = theme.white if dark_bg else theme.primary
    text_box(slide, left, top, width, height, theme.wordmark_text,
             13 if height < Inches(0.5) else 18, bold=True, color=color, align=align,
             font_name=theme.font_head, letter_spacing=0.5)


# ---------------------------------------------------------------------------
# Styled bullet frame and chips
# ---------------------------------------------------------------------------

def bullet_frame(slide, theme: Theme, left, top, width, height, bullets: list[str],
                 font_size: float = 15, min_size: float = 12, marker: str = "—",
                 dark_bg: bool = False):
    """Bullets with an accent em-dash marker, airy spacing and 1.22 leading,
    fitted to the box (smaller font first, then fewer bullets).

    `dark_bg` must be set true when this frame is drawn over a dark/photo
    backdrop so the body text switches to a light colour.
    """
    size, bullets = fit_bullets(bullets, width, height, font_size, min_size,
                                marker=f"{marker}   ", line_spacing=1.22)
    tb = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    marker_color = theme.secondary if dark_bg else theme.secondary_text
    text_color = theme.white if dark_bg else theme.dark_text

    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else round(size * BULLET_GAP_EM))
        p.space_after  = Pt(0)
        p.line_spacing = 1.22

        r_mark = p.add_run()
        r_mark.text = f"{marker}   "
        r_mark.font.size = Pt(size)
        r_mark.font.color.rgb = marker_color
        r_mark.font.bold = True
        r_mark.font.name = theme.font_body

        r_text = p.add_run()
        r_text.text = text
        r_text.font.size = Pt(size)
        r_text.font.color.rgb = text_color
        r_text.font.name = theme.font_body

    return tb


def lead_and_bullets(slide, theme: Theme, left, top, width, height, bullets: list[str],
                     lead: str = "", font_size: float = 16):
    """Optional intro paragraph, then bullets underneath. The lead gets at
    most ~40% of the box; the bullets fit into whatever is left."""
    if lead:
        lead_h = min(height * 0.4, text_height(lead.split("\n"), width, font_size - 1,
                                               line_spacing=1.3) + Inches(0.05))
        text_box(slide, left, top, width, lead_h, lead, font_size - 1, color=theme.dark_text,
                 font_name=theme.font_body, line_spacing=1.3, min_size=11)
        top, height = top + lead_h + Inches(0.15), height - lead_h - Inches(0.15)
    if bullets:
        bullet_frame(slide, theme, left, top, width, height, bullets, font_size=font_size)


def chip_row(slide, theme: Theme, left, top, items: list[str], font_size: int = 11,
             right=None, max_rows: int = 2):
    """Rounded tag chips, wrapping onto at most `max_rows` rows. Returns the
    y-coord just below the last row."""
    right = right if right is not None else W - MARGIN
    x = left
    chip_h = Inches(0.36)
    gap    = Inches(0.14)
    char_w = Pt(font_size) * 0.64
    row = 1

    for item in items[:10]:
        chip_w = int(char_w * len(item)) + int(Inches(0.4))
        if x + chip_w > right and x > left:
            if row == max_rows:
                break
            row += 1
            x = left
            top = top + chip_h + gap
        s = rect(slide, x, top, chip_w, chip_h, theme.light_bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = s.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = item
        r.font.size = Pt(font_size)
        r.font.color.rgb = theme.primary
        r.font.name = theme.font_body
        r.font.bold = True
        x += chip_w + gap

    return top + chip_h


# ---------------------------------------------------------------------------
# Slide layouts
# ---------------------------------------------------------------------------

def dual_accent_column(slide, theme: Theme):
    """Right-edge accent + secondary bars shared by the dark brand slides."""
    rect(slide, W - Inches(1.6), 0, Inches(0.14), H, theme.accent)
    rect(slide, W - Inches(1.34), 0, Inches(0.07), H, theme.secondary)


def slide_title_cover(prs, theme: Theme, title: str, subtitle: str, label: str = ""):
    slide = blank_slide(prs)
    rect(slide, 0, 0, W, H, theme.primary)
    dual_accent_column(slide, theme)
    logo(slide, theme, dark_bg=True, left=MARGIN, top=Inches(0.6), height=Inches(0.46))

    if label:
        rect(slide, MARGIN, Inches(2.4), Inches(0.34), Inches(0.05), theme.secondary)
        text_box(slide, MARGIN, Inches(2.55), W * 0.7, Inches(0.4),
                 label.upper(), 11, bold=True, color=theme.accent,
                 font_name=theme.font_body, letter_spacing=2.5)

    text_box(slide, MARGIN - Inches(0.03), Inches(3.05), W * 0.7, Inches(2.3),
             title, 40, bold=True, color=theme.white, line_spacing=1.04,
             font_name=theme.font_head, min_size=26)

    if subtitle:
        text_box(slide, MARGIN, Inches(5.55), W * 0.7, Inches(0.9),
                 subtitle, 15, color=theme.accent_dim, font_name=theme.font_body, min_size=11)
    return slide


def content_header(slide, theme: Theme, heading: str, eyebrow: str = ""):
    """Light, premium content header: eyebrow + big dark title + accent rule.

    Returns the y-coord where content can start.
    """
    rect(slide, 0, 0, W, H, theme.off_white)

    top = Inches(0.62)
    if eyebrow:
        rect(slide, MARGIN, top + Inches(0.03), Inches(0.18), Inches(0.16), theme.secondary)
        text_box(slide, MARGIN + Inches(0.32), top, W - MARGIN * 2, Inches(0.32),
                 eyebrow.upper(), 10, bold=True, color=theme.accent_text,
                 font_name=theme.font_body, letter_spacing=2.0)
        top = top + Inches(0.34)

    text_box(slide, MARGIN - Inches(0.02), top, W - MARGIN * 2, Inches(0.72),
             heading, 26, bold=True, color=theme.primary, line_spacing=1.0,
             font_name=theme.font_head, min_size=17, anchor=MSO_ANCHOR.BOTTOM)

    rule_y = top + Inches(0.74)
    rect(slide, MARGIN, rule_y, Inches(0.62), Inches(0.045), theme.accent)
    rect(slide, MARGIN + Inches(0.62), rule_y, Inches(0.28), Inches(0.045), theme.secondary)
    rect(slide, MARGIN + Inches(0.9), rule_y + Inches(0.016), W - MARGIN * 2 - Inches(0.9),
         Inches(0.013), theme.hairline)

    # Footer: brand mark + small wordmark text (recurring brand cue)
    favicon(slide, theme, W - Inches(2.05), H - Inches(0.56), size=Inches(0.3))
    text_box(slide, W - Inches(1.65), H - Inches(0.54), Inches(1.3), Inches(0.3),
             theme.wordmark_text, 9, bold=True, color=theme.mid_gray,
             font_name=theme.font_head)

    return rule_y + Inches(0.32)


def slide_cover_with_image(prs, theme: Theme, title: str, subtitle: str,
                           label: str, image_bytes: bytes | None):
    """Cover with a photo on the right; without one, the plain title cover."""
    if not image_bytes:
        return slide_title_cover(prs, theme, title, subtitle, label=label)

    slide = blank_slide(prs)
    embed_image(slide, image_bytes, 0, 0, W, H, theme)
    # Solid panel behind the text (left 62%) for legibility; the right 38%
    # stays photo under only a light scrim.
    rect(slide, 0, 0, W * 0.62, H, theme.primary)
    scrim(slide, W * 0.62, 0, W * 0.38, H, theme.primary_dark, alpha_pct=22)
    rect(slide, W * 0.62 - Inches(0.06), 0, Inches(0.06), H, theme.accent)
    rect(slide, 0, H - Inches(0.22), W, Inches(0.22), theme.accent)
    rect(slide, 0, 0, STRIP_W, H, theme.accent)

    logo(slide, theme, dark_bg=True, left=MARGIN, top=Inches(0.6), height=Inches(0.46))
    if label:
        text_box(slide, MARGIN, Inches(2.2), W * 0.55, Inches(0.38),
                 label.upper(), 11, bold=True, color=theme.accent,
                 font_name=theme.font_body, letter_spacing=2.5)
    text_box(slide, MARGIN - Inches(0.03), Inches(2.65), W * 0.55, Inches(2.2),
             title, 36, bold=True, color=theme.white, line_spacing=1.04,
             font_name=theme.font_head, min_size=24)
    if subtitle:
        text_box(slide, MARGIN, Inches(5.0), W * 0.55, Inches(0.9),
                 subtitle, 15, color=theme.accent_dim, font_name=theme.font_body, min_size=11)
    return slide


def slide_section_divider(prs, theme: Theme, label: str, image_bytes: bytes | None = None,
                          eyebrow: str = "Section"):
    """Full-bleed dark divider with an eyebrow and a big section label. With
    `image_bytes` the image becomes a background under a dark scrim."""
    slide = blank_slide(prs)
    if image_bytes:
        embed_image(slide, image_bytes, 0, 0, W, H, theme)
        scrim(slide, 0, 0, W, H, theme.primary, alpha_pct=72)
    else:
        rect(slide, 0, 0, W, H, theme.primary)

    rect(slide, MARGIN, Inches(2.7), Inches(0.6), Inches(0.05), theme.accent)
    rect(slide, MARGIN + Inches(0.6), Inches(2.7), Inches(0.3), Inches(0.05), theme.secondary)
    text_box(slide, MARGIN, Inches(2.9), W - MARGIN * 2, Inches(0.4),
             eyebrow.upper(), 11, bold=True, color=theme.accent,
             font_name=theme.font_body, letter_spacing=2.5)
    text_box(slide, MARGIN - Inches(0.03), Inches(3.35), W - Inches(2.4), Inches(2.0),
             label, 34, bold=True, color=theme.white, line_spacing=1.02,
             font_name=theme.font_head, min_size=24)

    dual_accent_column(slide, theme)
    favicon(slide, theme, MARGIN, H - Inches(0.7), size=Inches(0.4))
    return slide


def slide_bullets(prs, theme: Theme, heading: str, bullets: list[str], note: str = "",
                  eyebrow: str = "", reserve_bottom=0, lead: str = ""):
    """Heading + bullets. `reserve_bottom` keeps a band above the footer free
    (e.g. for a chip row)."""
    slide = blank_slide(prs)
    content_top = content_header(slide, theme, heading, eyebrow)
    top = content_top + Inches(0.15)
    lead_and_bullets(slide, theme, MARGIN, top, W - MARGIN * 2,
                     FOOTER_Y - top - reserve_bottom, bullets, lead, font_size=16)
    set_notes(slide, note)
    return slide


def slide_bullets_with_image(prs, theme: Theme, heading: str, bullets: list[str],
                             image_bytes: bytes | None, note: str = "",
                             side: str = "right", eyebrow: str = "", hint: str = "",
                             lead: str = ""):
    """Two-zone: bullets on one side, image (or a placeholder) on the other."""
    slide = blank_slide(prs)
    content_top = content_header(slide, theme, heading, eyebrow=eyebrow)

    img_w = W * 0.42
    gap   = Inches(0.4)
    img_t = content_top + Inches(0.1)
    img_h = FOOTER_Y - img_t - Inches(0.1)

    if side == "left":
        img_l  = MARGIN
        text_l = img_l + img_w + gap
        text_w = W - text_l - MARGIN
    else:
        img_l  = W - MARGIN - img_w
        text_l = MARGIN
        text_w = img_l - gap - MARGIN

    top = content_top + Inches(0.15)
    lead_and_bullets(slide, theme, text_l, top, text_w, FOOTER_Y - top, bullets, lead,
                     font_size=15)

    if image_bytes:
        rect(slide, img_l - Inches(0.04), img_t - Inches(0.04),
             img_w + Inches(0.08), img_h + Inches(0.08), theme.accent)
        embed_image(slide, image_bytes, img_l, img_t, img_w, img_h, theme)
        set_notes(slide, note)
    else:
        image_placeholder(slide, theme, img_l, img_t, img_w, img_h, hint)
        set_notes(slide, note, placeholder_note(hint))
    return slide


def slide_image_full(prs, theme: Theme, heading: str, bullets: list[str],
                     image_bytes: bytes, note: str = ""):
    """Full-bleed image with a dark scrim band along the bottom carrying the
    heading + a couple of short bullets. Dramatic, used sparingly — only with
    a real image (see glaze/place.py)."""
    slide = blank_slide(prs)
    embed_image(slide, image_bytes, 0, 0, W, H, theme)

    band_h = Inches(2.7)
    scrim(slide, 0, H - band_h, W, band_h, theme.primary_dark, alpha_pct=78)
    rect(slide, 0, H - band_h, Inches(0.14), band_h, theme.accent)
    rect(slide, Inches(0.14), H - band_h, Inches(0.07), band_h, theme.secondary)

    text_box(slide, MARGIN, H - band_h + Inches(0.28), W - MARGIN * 2, Inches(0.7),
             heading, 26, bold=True, color=theme.white, font_name=theme.font_head,
             line_spacing=1.0, min_size=18)
    if bullets:
        bullet_frame(slide, theme, MARGIN, H - band_h + Inches(1.1),
                     W - MARGIN * 2, band_h - Inches(1.3), bullets[:3], font_size=14,
                     dark_bg=True)
    favicon(slide, theme, W - Inches(0.95), Inches(0.5), size=Inches(0.4))
    set_notes(slide, note)
    return slide


def slide_content(prs, theme: Theme, heading: str, body: str, note: str = "",
                  eyebrow: str = ""):
    slide = blank_slide(prs)
    content_top = content_header(slide, theme, heading, eyebrow)
    top = content_top + Inches(0.15)
    body_text(slide, theme, MARGIN, top, W - MARGIN * 2, FOOTER_Y - top, body, size=16)
    set_notes(slide, note)
    return slide


def slide_two_col(prs, theme: Theme, heading: str,
                  left_head: str, left_body: str,
                  right_head: str, right_body: str, eyebrow: str = ""):
    slide = blank_slide(prs)
    content_top = content_header(slide, theme, heading, eyebrow)

    col_w = (W - MARGIN * 3) / 2
    body_top = content_top + Inches(0.46)
    body_h = FOOTER_Y - body_top

    rect(slide, MARGIN + col_w + MARGIN * 0.5 - Inches(0.01),
         content_top + Inches(0.1), Inches(0.02), FOOTER_Y - content_top - Inches(0.1),
         theme.hairline)

    for x, head, body in ((MARGIN, left_head, left_body),
                          (MARGIN * 2 + col_w, right_head, right_body)):
        text_box(slide, x, content_top, col_w, Inches(0.34),
                 head.upper(), 11, bold=True, color=theme.accent_text,
                 font_name=theme.font_body, letter_spacing=1.5)
        body_text(slide, theme, x, body_top, col_w, body_h, body, size=15)
    return slide


def slide_tech_stack(prs, theme: Theme, heading: str, body: str, chips: list[str],
                     eyebrow: str = ""):
    """Prose on the architecture, then the technologies as chips."""
    slide = blank_slide(prs)
    content_top = content_header(slide, theme, heading, eyebrow)
    top = content_top + Inches(0.1)
    chips_h = Inches(0.9) if chips else 0
    body_text(slide, theme, MARGIN, top, W - MARGIN * 2,
              FOOTER_Y - top - chips_h - Inches(0.3), body, size=16)
    if chips:
        chip_row(slide, theme, MARGIN, FOOTER_Y - chips_h, chips)
    return slide


def slide_kpis(prs, theme: Theme, heading: str, metrics: list[dict], eyebrow: str = "",
               note: str = ""):
    """KPI tiles: big secondary-colour numbers with captions. `metrics` = [{value,label}]."""
    slide = blank_slide(prs)
    content_top = content_header(slide, theme, heading, eyebrow)
    kpi_tiles(slide, theme, MARGIN, content_top + Inches(0.4), W - MARGIN * 2,
              Inches(2.2), metrics[:4])
    set_notes(slide, note)
    return slide


def kpi_tiles(slide, theme: Theme, left, top, width, height, metrics: list[dict],
              value_size: float = 46):
    n = len(metrics)
    if n == 0:
        return
    gap = Inches(0.3)
    tile_w = (width - gap * (n - 1)) / n
    for i, m in enumerate(metrics):
        x = left + i * (tile_w + gap)
        rect(slide, x, top, tile_w, height, theme.white)
        rect(slide, x, top, tile_w, Inches(0.08), theme.secondary)
        for bx, by, bw, bh in ((x, top, Inches(0.02), height),
                               (x + tile_w - Inches(0.02), top, Inches(0.02), height),
                               (x, top + height - Inches(0.02), tile_w, Inches(0.02))):
            rect(slide, bx, by, bw, bh, theme.hairline)
        text_box(slide, x + Inches(0.1), top + height * 0.18, tile_w - Inches(0.2), height * 0.45,
                 str(m.get("value", "")), value_size, bold=True, color=theme.secondary_text,
                 align=PP_ALIGN.CENTER, font_name=theme.font_head, line_spacing=1.0,
                 min_size=value_size * 0.55, anchor=MSO_ANCHOR.MIDDLE)
        text_box(slide, x + Inches(0.15), top + height * 0.66, tile_w - Inches(0.3),
                 height * 0.3, str(m.get("label", "")), 12, color=theme.mid_gray,
                 align=PP_ALIGN.CENTER, font_name=theme.font_body, line_spacing=1.15,
                 min_size=9)


def slide_chart(prs, theme: Theme, heading: str, chart_spec: dict,
                bullets: list[str] | None = None, note: str = "", eyebrow: str = ""):
    """Slide with a native editable chart (left) and optional bullets (right)."""
    from .charts import add_chart

    slide = blank_slide(prs)
    content_top = content_header(slide, theme, heading, eyebrow)

    has_bullets = bool(bullets)
    chart_w = (W * 0.58) if has_bullets else (W - MARGIN * 2)
    chart_l = MARGIN
    chart_t = content_top + Inches(0.2)

    title = chart_spec.get("title")
    if title:
        text_box(slide, chart_l, content_top, chart_w, Inches(0.3),
                 str(title).upper(), 10, bold=True, color=theme.accent_text,
                 font_name=theme.font_body, letter_spacing=1.5)
        chart_t = content_top + Inches(0.42)
    chart_h = FOOTER_Y - chart_t

    if add_chart(slide, chart_spec, chart_l, chart_t, chart_w, chart_h, theme) is None:
        rect(slide, chart_l, chart_t, chart_w, chart_h, theme.light_bg)

    if has_bullets:
        bx = MARGIN + chart_w + Inches(0.4)
        bullet_frame(slide, theme, bx, chart_t, W - bx - MARGIN, chart_h, bullets, font_size=14)
    set_notes(slide, note)
    return slide


def slide_statement(prs, theme: Theme, headline: str, sub: str = "",
                    image_bytes: bytes | None = None, note: str = ""):
    """One bold idea, centered. Dark slide; optional dimmed photo background."""
    slide = blank_slide(prs)
    if image_bytes:
        embed_image(slide, image_bytes, 0, 0, W, H, theme)
        scrim(slide, 0, 0, W, H, theme.primary, alpha_pct=62)
    else:
        rect(slide, 0, 0, W, H, theme.primary)

    rect(slide, W / 2 - Inches(0.45), Inches(2.35), Inches(0.6), Inches(0.06), theme.secondary)
    rect(slide, W / 2 + Inches(0.15), Inches(2.35), Inches(0.3), Inches(0.06), theme.accent)
    text_box(slide, Inches(1.2), Inches(2.7), W - Inches(2.4), Inches(1.9),
             headline, 36, bold=True, color=theme.white, align=PP_ALIGN.CENTER,
             line_spacing=1.05, font_name=theme.font_head, min_size=24)
    if sub:
        text_box(slide, Inches(1.8), Inches(4.7), W - Inches(3.6), Inches(1.4),
                 sub, 17, color=theme.accent_dim, align=PP_ALIGN.CENTER,
                 font_name=theme.font_body, line_spacing=1.25, min_size=12)
    favicon(slide, theme, W - Inches(0.95), H - Inches(0.78), size=Inches(0.4))
    set_notes(slide, note)
    return slide


def slide_quote(prs, theme: Theme, quote: str, attribution: str = "", note: str = ""):
    """A single pulled quote on a light slide with a big secondary-coloured mark."""
    slide = blank_slide(prs)
    rect(slide, 0, 0, W, H, theme.off_white)
    rect(slide, MARGIN, Inches(1.6), Inches(0.12), H - Inches(3.2), theme.secondary)
    text_box(slide, MARGIN + Inches(0.5), Inches(1.4), Inches(1.5), Inches(1.2),
             "“", 90, bold=True, color=theme.hairline)
    text_box(slide, MARGIN + Inches(0.55), Inches(2.5), W - MARGIN * 2 - Inches(1.0),
             Inches(2.6), quote, 26, color=theme.primary, font_name=theme.font_head,
             line_spacing=1.2, min_size=16)
    if attribution:
        text_box(slide, MARGIN + Inches(0.55), H - Inches(1.5),
                 W - MARGIN * 2 - Inches(1.0), Inches(0.5),
                 attribution.upper(), 11, bold=True, color=theme.accent_text,
                 font_name=theme.font_body, letter_spacing=1.5)
    favicon(slide, theme, W - Inches(0.95), H - Inches(0.78), size=Inches(0.4))
    set_notes(slide, note)
    return slide


def slide_diagram(prs, theme: Theme, heading: str, diagram_spec: dict,
                  caption: str = "", note: str = "", eyebrow: str = "Architecture"):
    """Slide with a native editable architecture/flow diagram."""
    from .diagrams import add_diagram

    slide = blank_slide(prs)
    content_top = content_header(slide, theme, heading, eyebrow)

    diag_t = content_top + Inches(0.2)
    diag_h = FOOTER_Y - diag_t - (Inches(0.5) if caption else 0)
    if not add_diagram(slide, diagram_spec, MARGIN, diag_t, W - MARGIN * 2, diag_h, theme):
        rect(slide, MARGIN, diag_t, W - MARGIN * 2, diag_h, theme.light_bg)

    if caption:
        # Kept clear of the recurring footer brand mark/wordmark.
        text_box(slide, MARGIN, H - Inches(1.05), W - MARGIN * 2, Inches(0.45),
                 caption, 12, color=theme.mid_gray, align=PP_ALIGN.CENTER,
                 font_name=theme.font_body, line_spacing=1.15, min_size=10)
    set_notes(slide, note)
    return slide


def slide_closing(prs, theme: Theme, tagline: str | None = None):
    """Brand sign-off: the tagline if the theme has one (otherwise the
    company name), website, logo — and the small Deck Mixer credit."""
    slide = blank_slide(prs)
    rect(slide, 0, 0, W, H, theme.primary)
    dual_accent_column(slide, theme)
    center_w = W - Inches(1.6)                  # centre in the area left of the accent column

    rect(slide, center_w / 2 - Inches(0.45), Inches(2.55), Inches(0.6), Inches(0.05), theme.accent)
    rect(slide, center_w / 2 + Inches(0.15), Inches(2.55), Inches(0.3), Inches(0.05), theme.secondary)

    headline = tagline or theme.tagline or theme.company_name
    text_box(slide, Inches(0.8), Inches(2.9), center_w - Inches(1.6), Inches(1.3),
             headline, 42, bold=True, color=theme.white, align=PP_ALIGN.CENTER,
             font_name=theme.font_head, line_spacing=1.05, min_size=26,
             anchor=MSO_ANCHOR.MIDDLE)
    if theme.website:
        text_box(slide, 0, Inches(4.35), center_w, Inches(0.5),
                 theme.website, 14, bold=True, color=theme.accent,
                 align=PP_ALIGN.CENTER, font_name=theme.font_body, letter_spacing=1.5)

    # The logo sits below — unless the headline already is the bare name.
    if headline != theme.company_name or (theme.logo_path and theme.logo_path.exists()):
        logo(slide, theme, dark_bg=True, left=center_w / 2 - Inches(1.5), top=Inches(5.05),
             height=Inches(0.55), width=Inches(3.0), align=PP_ALIGN.CENTER)

    if theme.show_credit:
        text_box(slide, 0, H - Inches(0.62), center_w, Inches(0.32), CREDIT, 9,
                 color=theme.accent_dim, align=PP_ALIGN.CENTER, font_name=theme.font_body,
                 letter_spacing=0.5)
    return slide
