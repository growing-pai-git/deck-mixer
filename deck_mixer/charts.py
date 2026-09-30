"""Native (editable) PowerPoint charts for the deck mixer.

Charts are real Pptx chart objects — fully editable in PowerPoint, not
rasterized images. Driven by an optional `chart` block in a case's YAML
frontmatter, e.g.:

    chart:
      type: column            # column | bar | line | donut
      title: Automation rate
      categories: [Manual, Automated]
      series:
        Order emails: [20, 80]

and optional `metrics` KPI tiles:

    metrics:
      - { value: "80%",   label: "order emails automated" }
      - { value: "2000+", label: "documents per month" }
"""

from __future__ import annotations

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.util import Pt

from .theme import DEFAULT_THEME, Theme

_TYPE_MAP = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar":    XL_CHART_TYPE.BAR_CLUSTERED,
    "line":   XL_CHART_TYPE.LINE_MARKERS,
    "donut":  XL_CHART_TYPE.DOUGHNUT,
    "pie":    XL_CHART_TYPE.PIE,
}


def _series_colors(theme: Theme) -> list[RGBColor]:
    # Secondary leads (it's the brand's "pop" accent), primary accent follows.
    return [theme.secondary, theme.accent, theme.primary, theme.mid_gray]


def add_chart(slide, spec: dict, left, top, width, height,
              theme: Theme = DEFAULT_THEME) -> object | None:
    """Add a native editable chart from a `chart` spec dict.

    Returns the chart graphic frame, or None if the spec is malformed.
    """
    try:
        ctype = _TYPE_MAP.get(str(spec.get("type", "column")).lower(),
                              XL_CHART_TYPE.COLUMN_CLUSTERED)
        categories = spec.get("categories") or []
        series = spec.get("series") or {}
        if not categories or not series:
            return None

        data = CategoryChartData()
        data.categories = [str(c) for c in categories]
        for name, values in series.items():
            data.add_series(str(name), tuple(float(v) for v in values))

        gframe = slide.shapes.add_chart(
            ctype, int(left), int(top), int(width), int(height), data
        )
        chart = gframe.chart
        _style_chart(chart, ctype, n_series=len(series), theme=theme)
        return gframe
    except Exception:
        return None


def _style_chart(chart, ctype, n_series: int, theme: Theme):
    """Apply clean brand styling: no clutter, theme colours, data labels."""
    chart.has_title = False

    # Legend only when more than one series
    if n_series > 1:
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)
        chart.legend.font.name = theme.font_body
    else:
        chart.has_legend = False

    is_round = ctype in (XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.PIE)
    colors = _series_colors(theme)

    for i, plot_series in enumerate(chart.series):
        color = colors[i % len(colors)]
        if is_round:
            # Colour each point of the single series distinctly
            for j, point in enumerate(plot_series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = colors[j % len(colors)]
        else:
            plot_series.format.fill.solid()
            plot_series.format.fill.fore_color.rgb = color
            plot_series.format.line.fill.background()

    # Data labels
    try:
        plot = chart.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(11)
        dl.font.bold = True
        dl.font.name = theme.font_body
        dl.number_format = "General"
        dl.number_format_is_linked = False
        if not is_round:
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
    except Exception:
        pass

    # De-clutter axes for bar/column/line
    try:
        if not is_round:
            cat_ax = chart.category_axis
            cat_ax.tick_labels.font.size = Pt(11)
            cat_ax.tick_labels.font.name = theme.font_body
            cat_ax.format.line.color.rgb = theme.hairline
            val_ax = chart.value_axis
            val_ax.visible = False
            val_ax.has_major_gridlines = False
    except Exception:
        pass
