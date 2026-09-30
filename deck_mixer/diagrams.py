"""Native (editable) architecture / flow diagrams for the deck mixer.

Renders branded boxes + connectors directly as PowerPoint shapes — editable
in PPT, no external dependency. This is the "draw.io" piece, kept
self-contained so the builder needs no runtime services.

Driven by an optional `diagram` block in a case's YAML frontmatter:

    diagram:
      type: flow                      # flow (left->right) | stack (top->bottom)
      nodes:
        - { id: src,  label: "Mailbox" }
        - { id: ext,  label: "Document\\nIntelligence", accent: true }
        - { id: llm,  label: "Model API" }
        - { id: ui,   label: "Validation UI" }
        - { id: erp,  label: "ERP" }
      edges:
        - [src, ext]
        - [ext, llm]
        - [llm, ui]
        - [ui, erp]

`accent: true` highlights a node in the theme's secondary colour (the hero
component).
"""

from __future__ import annotations

from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from .theme import DEFAULT_THEME, Theme

_ROUNDED = 5  # MSO rounded rectangle autoshape id


def add_diagram(slide, spec: dict, left, top, width, height,
                theme: Theme = DEFAULT_THEME) -> bool:
    """Render a flow/stack diagram into the given rectangle. Returns True on
    success, False if the spec is unusable (caller can draw a fallback)."""
    try:
        nodes = spec.get("nodes") or []
        edges = spec.get("edges") or []
        if not nodes:
            return False
        dtype = str(spec.get("type", "flow")).lower()
        horizontal = dtype != "stack"

        n = len(nodes)
        boxes = {}

        if horizontal:
            gap = Inches(0.45)
            total_gap = gap * (n - 1)
            box_w = (width - total_gap) / n
            box_h = min(Inches(1.5), height * 0.55)
            cy = top + (height - box_h) / 2
            for i, node in enumerate(nodes):
                x = left + i * (box_w + gap)
                boxes[node.get("id", i)] = _node(slide, node, x, cy, box_w, box_h, theme)
        else:
            gap = Inches(0.32)
            total_gap = gap * (n - 1)
            box_h = (height - total_gap) / n
            box_w = min(Inches(4.2), width * 0.7)
            cx = left + (width - box_w) / 2
            for i, node in enumerate(nodes):
                y = top + i * (box_h + gap)
                boxes[node.get("id", i)] = _node(slide, node, cx, y, box_w, box_h, theme)

        # Edges: explicit list, else chain consecutive nodes
        if not edges:
            ids = [node.get("id", i) for i, node in enumerate(nodes)]
            edges = list(zip(ids, ids[1:]))
        for edge in edges:
            a, b = (edge[0], edge[1]) if isinstance(edge, (list, tuple)) else (None, None)
            if a in boxes and b in boxes:
                _connect(slide, boxes[a], boxes[b], horizontal, theme)

        return True
    except Exception:
        return False


def _node(slide, node: dict, x, y, w, h, theme: Theme):
    accent = bool(node.get("accent"))
    accent_fill = theme.secondary_text
    shape = slide.shapes.add_shape(_ROUNDED, int(x), int(y), int(w), int(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = accent_fill if accent else theme.white
    shape.line.color.rgb = accent_fill if accent else theme.hairline
    shape.line.width = Pt(1.5 if accent else 1.0)
    shape.shadow.inherit = False

    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(6); tf.margin_right = Pt(6)
    tf.margin_top = Pt(4); tf.margin_bottom = Pt(4)
    label = str(node.get("label", "")).replace("\\n", "\n")
    for i, line in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = line
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.name = theme.font_body
        r.font.color.rgb = theme.white if accent else theme.primary
    return shape


def _connect(slide, a, b, horizontal: bool, theme: Theme):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    conn.begin_connect(a, 3 if horizontal else 2)   # right / bottom of A
    conn.end_connect(b, 1 if horizontal else 0)     # left / top of B
    conn.line.color.rgb = theme.accent
    conn.line.width = Pt(2.0)
    # Arrowhead at the end
    try:
        from pptx.oxml.ns import qn
        ln = conn.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    except Exception:
        pass
    return conn
