"""Quality bar for generated decks: content reaches the slides (not just the
notes), text fits its box, confidential names never show, the file carries
proper metadata, and a zero-key build never touches the network."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Pt

from deck_mixer.textfit import BULLET_GAP_EM, text_height

SAMPLE_LIBRARY = Path(__file__).parent.parent / "examples" / "sample-library"

PLAN_MD = """## Kickoff
We start with a two-day discovery workshop.

## Phase 1 - Discovery
We interview the operations team, map the current intake process and collect
a representative sample of documents.

- Stakeholder interviews with operations and IT
- Process mapping of the current intake flow
- Baseline metrics: volume, error rate, handling time

## Phase 2 - Build
- Extraction pipeline with human review queue
- Integration with the existing case management system
"""


@pytest.fixture(autouse=True)
def _offline(monkeypatch):
    """No keys, no config — and any network call fails the test."""
    for var in (
        "ANTHROPIC_API_KEY", "OPENAI_API_KEY", "GEMINI_API_KEY",
        "STABILITY_API_KEY", "TOGETHER_API_KEY", "FAL_KEY", "REPLICATE_API_TOKEN",
        "PEXELS_API_KEY", "UNSPLASH_ACCESS_KEY",
    ):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setattr("deck_mixer.config.load_config", lambda: {})

    def no_network(*a, **kw):
        raise AssertionError("zero-key build tried to reach the network")
    monkeypatch.setattr("urllib.request.urlopen", no_network)


@pytest.fixture(scope="module")
def cases():
    from deck_mixer.kb import filter_cases, load_catalog
    return filter_cases(load_catalog(SAMPLE_LIBRARY))


def _decks(tmp_path, cases) -> dict[str, Presentation]:
    from deck_mixer.builder import (
        build_capabilities_deck, build_plan_deck, build_reference_deck, build_tender_deck,
    )
    paths = {
        "reference": build_reference_deck(cases, SAMPLE_LIBRARY, tmp_path / "r.pptx"),
        "capabilities": build_capabilities_deck(cases, SAMPLE_LIBRARY, tmp_path / "c.pptx"),
        "tender": build_tender_deck(cases, SAMPLE_LIBRARY, tmp_path / "t.pptx", brief="A brief."),
        "plan": build_plan_deck(title="Sample Plan", content_md=PLAN_MD,
                                output_path=tmp_path / "p.pptx"),
    }
    return {k: Presentation(v) for k, v in paths.items()}


@pytest.fixture
def decks(tmp_path, cases):
    return _decks(tmp_path, cases)


def _slide_text(prs) -> str:
    """All visible text (shapes and groups), excluding speaker notes."""
    out = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:                  # group
                walk(sh.shapes)
            elif sh.has_text_frame:
                out.append(sh.text_frame.text)
    for slide in prs.slides:
        walk(slide.shapes)
    return " ".join(" ".join(out).split())


def _first_words(text: str, n: int = 6) -> str:
    from deck_mixer.kb import strip_markdown
    return " ".join(strip_markdown(text).split()[:n])


def test_reference_sections_are_on_slides(decks, cases):
    from deck_mixer.kb import extract_section, read_case
    text = _slide_text(decks["reference"])
    for case in cases:
        content = read_case(SAMPLE_LIBRARY, case)
        for n in (1, 2, 3, 5, 6, 7, 10):
            section = extract_section(content, n)
            assert _first_words(section) in text, f"{case['case_id']} §{n} missing from slides"


def test_tender_and_plan_content_on_slides(decks, cases):
    from deck_mixer.kb import extract_section, read_case
    tender = _slide_text(decks["tender"])
    for case in cases:
        summary = extract_section(read_case(SAMPLE_LIBRARY, case), 1)
        assert _first_words(summary, 4) in tender

    plan = _slide_text(decks["plan"])
    for needle in ("two-day discovery workshop", "We interview the operations team",
                   "Stakeholder interviews", "Extraction pipeline"):
        assert needle in plan


def test_confidential_client_never_on_slides(decks):
    for name, prs in decks.items():
        assert "BlueCrate" not in _slide_text(prs), f"confidential name leaked in {name}"
        for slide in prs.slides:
            if slide.has_notes_slide:
                assert "BlueCrate" not in slide.notes_slide.notes_text_frame.text


def test_no_hard_wrapped_lines(decks):
    """Paragraphs must not break mid-sentence (hard-wrapped markdown)."""
    for prs in decks.values():
        for slide in prs.slides:
            for sh in slide.shapes:
                if not sh.has_text_frame:
                    continue
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        assert "\n" not in r.text
                        assert not re.search(r"[a-z,]\n[a-z]", r.text)


def test_text_fits_its_box(decks):
    for name, prs in decks.items():
        for i, slide in enumerate(prs.slides, 1):
            for sh in slide.shapes:
                if not sh.has_text_frame or not sh.text_frame.text.strip():
                    continue
                paras = [p for p in sh.text_frame.paragraphs if p.runs]
                size = max((r.font.size or Pt(18)).pt for p in paras for r in p.runs)
                gap = max((p.space_before.pt if p.space_before else 0) for p in paras)
                need = text_height([p.text for p in paras], sh.width, size,
                                   line_spacing=paras[0].line_spacing or 1.0,
                                   para_gap_pt=gap if gap else 0,
                                   bold=bool(paras[0].runs[0].font.bold))
                if sh.text_frame.word_wrap is False:
                    continue
                assert need <= sh.height * 1.05 or size <= 8, \
                    f"{name} slide {i}: '{sh.text_frame.text[:40]}' overflows"
    assert BULLET_GAP_EM > 0


def test_metadata_and_credit(decks):
    from deck_mixer.slides import CREDIT
    for name, prs in decks.items():
        cp = prs.core_properties
        assert cp.author and cp.author != "Steve Canny", name
        assert cp.title, name
        assert cp.created.year >= 2026, name
        assert CREDIT in _slide_text(prs), f"credit missing from {name}"


def test_credit_can_be_switched_off(tmp_path, cases):
    from deck_mixer.builder import build_reference_deck
    from deck_mixer.slides import CREDIT
    from deck_mixer.theme import load_theme

    (tmp_path / "theme.yaml").write_text("company_name: Acme\ncredit: false\n")
    theme = load_theme(theme_path=str(tmp_path / "theme.yaml"))
    out = build_reference_deck(cases[:1], SAMPLE_LIBRARY, tmp_path / "x.pptx", theme=theme)
    assert CREDIT not in _slide_text(Presentation(out))


def test_zero_keys_gives_placeholders_not_photos(decks):
    from deck_mixer.slides import PLACEHOLDER_NAME
    for name, prs in decks.items():
        for slide in prs.slides:
            assert not any(sh.shape_type == 13 for sh in slide.shapes), \
                f"{name}: picture in a zero-key build"
    names = [sh.name for s in decks["plan"].slides for sh in s.shapes]
    assert PLACEHOLDER_NAME in names


def test_custom_template(tmp_path, cases):
    """A corporate .pptx template: its masters are used, no stray placeholders."""
    from pptx.util import Inches
    from deck_mixer.builder import build_reference_deck

    tpl = Presentation()
    tpl.slide_width, tpl.slide_height = Inches(13.333), Inches(7.5)
    tpl.save(tmp_path / "tpl.pptx")
    out = build_reference_deck(cases[:1], SAMPLE_LIBRARY, tmp_path / "t.pptx",
                               template_path=str(tmp_path / "tpl.pptx"))
    for slide in Presentation(out).slides:
        assert not list(slide.placeholders)

    tpl43 = Presentation()                            # default 4:3
    tpl43.save(tmp_path / "43.pptx")
    with pytest.raises(ValueError, match="16:9"):
        build_reference_deck(cases[:1], SAMPLE_LIBRARY, tmp_path / "u.pptx",
                             template_path=str(tmp_path / "43.pptx"))


def test_aliases_redacted_and_catalog_refreshes(tmp_path):
    """An alias in frontmatter is redacted too, and editing case.md is picked
    up without a manual `library validate`."""
    import os
    import shutil
    from deck_mixer.builder import build_tender_deck
    from deck_mixer.kb import filter_cases, load_catalog

    lib = tmp_path / "lib"
    shutil.copytree(SAMPLE_LIBRARY, lib)
    case_md = lib / "cases" / "bluecrate-route-optimization" / "case.md"
    text = case_md.read_text(encoding="utf-8")
    text = text.replace("client: BlueCrate Logistics", "client: BlueCrate Logistics\naliases: [BCL]")
    text = text.replace("## 1. Summary\n", "## 1. Summary\n\nBCL ships parcels.\n", 1)
    case_md.write_text(text, encoding="utf-8")
    catalog_mtime = (lib / "catalog.json").stat().st_mtime
    os.utime(case_md, (catalog_mtime + 10, catalog_mtime + 10))

    cases = filter_cases(load_catalog(lib), case_ids=["sample-2026-bluecrate-route-optimization"])
    assert cases[0]["aliases"] == ["BCL"]
    prs = Presentation(build_tender_deck(cases, lib, tmp_path / "t.pptx"))
    text = _slide_text(prs)
    assert "BCL" not in text and "BlueCrate" not in text
    assert "The client ships parcels" in text


def _fake_gemini(monkeypatch, respond):
    """Route the Gemini provider through `respond(system, user) -> dict` instead
    of the network, with a (fake) key set so the AI planning path is taken."""
    import json
    from deck_mixer import llm

    monkeypatch.setenv("GEMINI_API_KEY", "test-key")
    monkeypatch.delenv("LLM_PROVIDER", raising=False)
    calls = []

    def fake_http(url, *, data, headers, timeout=60):
        assert "generativelanguage.googleapis.com" in url
        assert headers["X-goog-api-key"] == "test-key"
        body = json.loads(data)
        system = body["systemInstruction"]["parts"][0]["text"]
        user = body["contents"][0]["parts"][0]["text"]
        calls.append(user)
        answer = json.dumps(respond(system, user))
        return json.dumps({"candidates": [{"content": {"parts": [{"text": answer}]}}]}).encode()

    monkeypatch.setattr(llm, "_http", fake_http)
    return calls


def test_ai_plan_drives_the_slides(tmp_path, monkeypatch):
    from deck_mixer.builder import build_plan_deck

    def respond(system, user):
        if not (m := re.search(r"EXACTLY (\d+) slide plans", system)):
            return {"mood": ["calm"], "subject_bias": "harbour cranes"}   # style call
        n = int(m.group(1))
        layouts = ["statement", "bullets", "image_right"]
        return {"slides": [
            {"layout": layouts[i % 3], "headline": f"Planned headline {i}",
             "bullets": [f"Planned bullet {i}"], "speaker_notes": f"Planned note {i}",
             "visual": {"want": False, "role": "split", "medium": "generate", "brief": ""}}
            for i in range(n)]}

    calls = _fake_gemini(monkeypatch, respond)
    out = build_plan_deck(title="AI Plan", content_md=PLAN_MD, enrich=False,
                          output_path=tmp_path / "ai.pptx")
    plan_calls = [c for c in calls if c.startswith("Plan a deck")]
    assert len(plan_calls) == 1 and "Phase 1 - Discovery" in plan_calls[0]
    text = _slide_text(Presentation(out))
    for i in range(3):
        assert f"Planned headline {i}" in text


@pytest.mark.parametrize("respond", [
    lambda system, user: {"slides": []},                  # wrong slide count
    lambda system, user: {"unexpected": True},            # malformed answer
    lambda system, user: (_ for _ in ()).throw(RuntimeError("HTTP 500: boom")),
])
def test_ai_plan_failure_falls_back_to_heuristics(tmp_path, monkeypatch, respond):
    from deck_mixer.builder import build_plan_deck

    _fake_gemini(monkeypatch, respond)
    out = build_plan_deck(title="AI Plan", content_md=PLAN_MD, enrich=False,
                          output_path=tmp_path / "fb.pptx")
    text = _slide_text(Presentation(out))
    assert "Stakeholder interviews" in text and "Planned headline" not in text
