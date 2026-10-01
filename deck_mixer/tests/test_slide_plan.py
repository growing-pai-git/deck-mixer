"""A slide plan supplied by the caller (Claude, via create_plan_deck's
slide_plan or the CLI's --plan) drives the deck through the same rendering
path as the internal planner, with per-slide fallback and a list of fixes."""

from __future__ import annotations

import importlib
import json
from pathlib import Path

import pytest
from pptx import Presentation

CONTENT = """## Why route planning matters
Drivers lose a day a week to bad routes.

## Phase 1 - Discovery
- Stakeholder interviews
- Route data audit
- Baseline metrics

## What we will deliver
- Route optimiser
- Dispatch dashboard
- Weekly reports
"""

GOOD_PLAN = [
    {"heading": "Why route planning matters", "layout": "statement",
     "headline": "A day a week lost to bad routes",
     "visual": {"want": True, "medium": "generate", "brief": "A delivery van at dawn"}},
    {"heading": "Phase 1 - Discovery", "layout": "bullets", "headline": "First, we listen",
     "bullets": ["Interviews with dispatch", "Audit of route data", "Baseline metrics"]},
    {"heading": "What we will deliver", "layout": "image_left", "headline": "Three things you get",
     "bullets": ["Route optimiser", "Dispatch dashboard", "Weekly reports"],
     "visual": {"want": True, "brief": "A dispatcher reviewing routes on a large screen"}},
]


@pytest.fixture(autouse=True)
def _offline_no_llm(monkeypatch):
    for var in ("GEMINI_API_KEY", "GOOGLE_API_KEY", "ANTHROPIC_API_KEY", "OPENAI_API_KEY",
                "STABILITY_API_KEY", "TOGETHER_API_KEY", "FAL_KEY", "REPLICATE_API_TOKEN",
                "PEXELS_API_KEY", "UNSPLASH_ACCESS_KEY", "LLM_PROVIDER"):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setattr("deck_mixer.config.load_config", lambda: {})

    def no_llm(*a, **kw):
        raise AssertionError("a supplied slide_plan must not call the planner LLM")
    monkeypatch.setattr("deck_mixer.llm.call_json", no_llm)

    def no_network(*a, **kw):
        raise AssertionError("tried to reach the network")
    monkeypatch.setattr("urllib.request.urlopen", no_network)


def _sections():
    from deck_mixer.builder import split_sections
    return split_sections(CONTENT)


def _text(path) -> str:
    return " ".join(sh.text_frame.text for s in Presentation(path).slides
                    for sh in s.shapes if sh.has_text_frame)


def test_valid_plan_passes_unchanged():
    from deck_mixer.planner import validate_plan
    plan, fixes = validate_plan(_sections(), GOOD_PLAN)
    assert fixes == []
    assert [p["layout"] for p in plan] == ["statement", "bullets", "image_left"]
    assert plan[2]["visual"] == {"want": True, "role": "split", "medium": "generate",
                                 "brief": "A dispatcher reviewing routes on a large screen"}


def test_each_problem_falls_back_for_that_slide_only():
    from deck_mixer.planner import validate_plan
    bad = [
        {"heading": "why ROUTE planning matters", "layout": "statement"},     # case-insensitive join
        {"heading": "Phase 1 - Discovery", "layout": "timeline"},             # unknown layout
        {"heading": "Budget", "layout": "bullets"},                           # no such heading
        {"layout": "bullets"},                                                # no heading
    ]                                                                         # last section missing
    plan, fixes = validate_plan(_sections(), bad)
    assert len(plan) == 3
    assert plan[0]["layout"] == "statement"
    assert plan[1]["layout"] != "timeline"
    text = "\n".join(fixes)
    assert "'timeline' is not one of" in text
    assert "'Budget' matches no" in text
    assert "entry 4: no heading" in text
    assert "slide 3 'What we will deliver': not in slide_plan" in text


def test_layouts_needing_data_and_long_statements_are_downgraded():
    from deck_mixer.planner import validate_plan
    plan, fixes = validate_plan(_sections(), [
        {"heading": "Phase 1 - Discovery", "layout": "statement"},    # 3 bullets: too long
        {"heading": "What we will deliver", "layout": "chart"},        # needs data
    ])
    assert plan[1]["layout"] == "bullets"
    assert plan[2]["layout"] != "chart"
    assert any("too much content for a statement" in f for f in fixes)


@pytest.mark.parametrize("junk", ["not a list", 42, [1, "x", None], {"slides": "nope"}])
def test_junk_input_never_crashes(junk):
    from deck_mixer.planner import validate_plan
    plan, fixes = validate_plan(_sections(), junk)
    assert len(plan) == 3 and fixes


def test_plan_drives_the_deck_without_the_llm(tmp_path):
    from deck_mixer.builder import build_plan_deck
    from deck_mixer.planner import validate_plan
    plan, _ = validate_plan(_sections(), {"slides": GOOD_PLAN})
    path = build_plan_deck(title="Routes", content_md=CONTENT, enrich=False,
                           output_path=tmp_path / "p.pptx", slide_plan=plan)
    text = _text(path)
    for headline in ("A day a week lost to bad routes", "First, we listen", "Three things you get"):
        assert headline in text


def test_rules_are_shared_by_planner_and_tool_description():
    from deck_mixer.planner import DESIGN_RULES, _build_system
    from deck_mixer.theme import DEFAULT_THEME
    assert DESIGN_RULES in _build_system(DEFAULT_THEME)
    server = importlib.import_module("deck_mixer.server")
    assert DESIGN_RULES in server.create_plan_deck.__doc__


@pytest.fixture
def server(tmp_path, monkeypatch):
    pytest.importorskip("mcp")
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path / "out"))
    monkeypatch.setenv("DECK_MIXER_BUNDLE", "1")          # skip first-run onboarding
    server = importlib.import_module("deck_mixer.server")
    monkeypatch.setattr(server, "OUTPUT_DIR", tmp_path / "out")
    (tmp_path / "out").mkdir()
    return server


def test_tool_plan_only_builds_nothing(server):
    result = json.loads(server.create_plan_deck(title="T", content_md=CONTENT,
                                                slide_plan=GOOD_PLAN, plan_only=True))
    assert [s["heading"] for s in result["slides"]] == [h for h, _ in _sections()]
    assert result["slides"][2]["headline"] == "Three things you get"
    assert result["fixes"] == []
    assert not list(Path(server.OUTPUT_DIR).iterdir())


def test_tool_reports_fixes_and_drops_the_text_key_notice(server):
    plan = GOOD_PLAN[:2] + [{"heading": "What we will deliver", "layout": "kpi"}]
    result = server.create_plan_deck(title="T", content_md=CONTENT, enrich=False, slide_plan=plan)
    assert "Created plan deck" in result
    assert "Plan adjustments:" in result and "'kpi' is not one of" in result
    assert "Text:" not in result


def test_tool_without_plan_tells_claude_it_can_design(server, monkeypatch):
    monkeypatch.setattr("deck_mixer.llm.call_json", lambda *a, **kw: None)   # no key: no answer
    result = server.create_plan_deck(title="T", content_md=CONTENT, enrich=False)
    assert "pass slide_plan" in result
