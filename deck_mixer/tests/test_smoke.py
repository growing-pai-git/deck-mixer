"""End-to-end smoke tests: every deck type must build from the sample
library and the default theme with zero API keys configured."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

SAMPLE_LIBRARY = Path(__file__).parent.parent / "examples" / "sample-library"


@pytest.fixture(autouse=True)
def _no_api_keys(monkeypatch):
    """Force the zero-config path regardless of the developer's own env/config."""
    for var in (
        "ANTHROPIC_API_KEY", "OPENAI_API_KEY", "GEMINI_API_KEY",
        "STABILITY_API_KEY", "TOGETHER_API_KEY", "FAL_KEY", "REPLICATE_API_TOKEN",
        "PEXELS_API_KEY", "UNSPLASH_ACCESS_KEY",
    ):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setattr("deck_mixer.config.load_config", lambda: {})


def _assert_valid_pptx(path: Path, min_slides: int = 1) -> None:
    assert path.exists()
    prs = Presentation(str(path))
    slides = list(prs.slides)
    assert len(slides) >= min_slides
    texts = [
        shape.text_frame.text.strip()
        for slide in slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    ]
    assert texts, "expected at least some slide text"
    assert not any(t.lower() == "none" for t in texts)


def test_library_loads():
    from deck_mixer.kb import load_catalog
    catalog = load_catalog(SAMPLE_LIBRARY)
    assert catalog["total_cases"] == 3


def test_build_reference_deck(tmp_path):
    from deck_mixer.builder import build_reference_deck
    from deck_mixer.kb import filter_cases, load_catalog

    cases = filter_cases(load_catalog(SAMPLE_LIBRARY), case_ids=["northwind-retail-support"])
    out = build_reference_deck(cases, SAMPLE_LIBRARY, tmp_path / "reference.pptx")
    _assert_valid_pptx(Path(out))


def test_build_capabilities_deck(tmp_path):
    from deck_mixer.builder import build_capabilities_deck
    from deck_mixer.kb import filter_cases, load_catalog

    cases = filter_cases(load_catalog(SAMPLE_LIBRARY))
    out = build_capabilities_deck(cases, SAMPLE_LIBRARY, tmp_path / "capabilities.pptx")
    _assert_valid_pptx(Path(out))


def test_build_tender_deck(tmp_path):
    from deck_mixer.builder import build_tender_deck
    from deck_mixer.kb import filter_cases, load_catalog

    cases = filter_cases(load_catalog(SAMPLE_LIBRARY), tender_tags=["human in the loop"])
    out = build_tender_deck(cases, SAMPLE_LIBRARY, tmp_path / "tender.pptx", brief="Sample tender.")
    _assert_valid_pptx(Path(out))


def test_build_plan_deck(tmp_path):
    from deck_mixer.builder import build_plan_deck

    content_md = "## Kickoff\nWe will start with discovery.\n\n## Phase 1 - Build\nBuild the core pipeline.\n"
    out = build_plan_deck(title="Sample Plan", content_md=content_md,
                          output_path=tmp_path / "plan.pptx", enrich=False)
    _assert_valid_pptx(Path(out))


def test_custom_theme_changes_colors(tmp_path):
    from deck_mixer.builder import build_reference_deck
    from deck_mixer.kb import filter_cases, load_catalog
    from deck_mixer.theme import DEFAULT_THEME, load_theme

    theme_file = tmp_path / "theme.yaml"
    theme_file.write_text(
        "colors:\n  primary: '2E1065'\n  accent: 'EC4899'\n  secondary: '22C55E'\n",
        encoding="utf-8",
    )
    theme = load_theme(theme_path=str(theme_file))
    assert str(theme.primary) != str(DEFAULT_THEME.primary)

    cases = filter_cases(load_catalog(SAMPLE_LIBRARY), case_ids=["northwind-retail-support"])
    out = build_reference_deck(cases, SAMPLE_LIBRARY, tmp_path / "themed.pptx", theme=theme)
    _assert_valid_pptx(Path(out))


def test_library_validate_rejects_broken_library(tmp_path):
    from deck_mixer.kb import load_catalog, rebuild_catalog

    lib = tmp_path / "broken"
    case_dir = lib / "cases" / "bad-case"
    case_dir.mkdir(parents=True)
    (case_dir / "case.md").write_text(
        "---\ncase_id: broken-1\ntitle: Broken\nclient: Test\n"
        "confidentiality: made_up_value\n---\n\n## 1. Summary\n\nOnly one section.\n",
        encoding="utf-8",
    )
    rebuild_catalog(lib)
    catalog = load_catalog(lib)
    case = catalog["cases"][0]
    assert len(case["sections"]) < 13
    assert case["confidentiality"] not in ("public", "anonymized", "confidential")
